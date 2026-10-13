"""
Multi-Sheet Excel Workbook Export Generator for GSTR-1 Draft.
Uses openpyxl to generate structured, styled, and numeric-formatted Excel reports.
"""
from pathlib import Path
from typing import Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.gst.gstr1_models import GSTR1Draft

# Style Constants
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="262626")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
REGULAR_FONT = Font(name="Calibri", size=11)
ALERT_FONT = Font(name="Calibri", size=11, bold=True, color="C00000")

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

CURRENCY_FORMAT = "₹#,##0.00"
PERCENT_FORMAT = "0.00%"
NUMBER_FORMAT = "#,##0.00"

def auto_fit_columns(ws, max_cols=20):
    for col in range(1, min(ws.max_column + 1, max_cols)):
        max_len = 0
        col_letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=col).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

def export_gstr1_excel(draft: GSTR1Draft, file_path: Union[str, Path]) -> Path:
    """
    Generates a professional multi-sheet Excel (.xlsx) workbook for the GSTR-1 draft.
    """
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # 1. SUMMARY SHEET
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title & Legal Disclaimer
    ws_sum["A1"] = "AI GST Invoice Intelligence — GSTR-1 Draft Summary"
    ws_sum["A1"].font = TITLE_FONT
    ws_sum["A2"] = "DRAFT STATUS: GSTR-1 DRAFT — NOT FILED (Preparation only. Not submitted to GST portal)"
    ws_sum["A2"].font = ALERT_FONT
    ws_sum["A3"] = f"Generated at: {draft.generated_at}"
    ws_sum["A3"].font = SUBTITLE_FONT

    # Section 1: Processing Counts
    ws_sum["A5"] = "1. INVOICE PROCESSING METRICS"
    ws_sum["A5"].font = SECTION_FONT

    metrics = [
        ("Total Source Invoices Evaluated", draft.source_invoice_count),
        ("Invoices Admitted into GSTR-1 Draft", draft.included_invoice_count),
        ("Invoices Excluded (Ineligible / Rejected)", draft.excluded_invoice_count),
        ("B2B Invoices (Table 4A - Registered Recipients)", draft.b2b_count),
        ("B2C Invoices (Table 7 - Unregistered Recipients)", draft.b2c_count),
    ]

    for idx, (label, val) in enumerate(metrics, start=6):
        ws_sum[f"A{idx}"] = label
        ws_sum[f"A{idx}"].font = REGULAR_FONT
        ws_sum[f"B{idx}"] = val
        ws_sum[f"B{idx}"].font = BOLD_FONT
        ws_sum[f"B{idx}"].alignment = Alignment(horizontal="right")

    # Section 2: Financial Aggregates
    ws_sum["A12"] = "2. FINANCIAL TOTALS (INR)"
    ws_sum["A12"].font = SECTION_FONT

    fin_totals = [
        ("Total Taxable Value", draft.total_taxable_value),
        ("Total CGST (Central Tax)", draft.total_cgst),
        ("Total SGST (State/UT Tax)", draft.total_sgst),
        ("Total IGST (Integrated Tax)", draft.total_igst),
        ("Total Grand Invoice Value", draft.total_invoice_value),
    ]

    for idx, (label, val) in enumerate(fin_totals, start=13):
        ws_sum[f"A{idx}"] = label
        ws_sum[f"A{idx}"].font = REGULAR_FONT
        ws_sum[f"B{idx}"] = float(val)
        ws_sum[f"B{idx}"].font = BOLD_FONT
        ws_sum[f"B{idx}"].number_format = NUMBER_FORMAT
        ws_sum[f"B{idx}"].alignment = Alignment(horizontal="right")

    auto_fit_columns(ws_sum)

    # -------------------------------------------------------------
    # 2. B2B_TABLE_4A SHEET
    # -------------------------------------------------------------
    ws_b2b = wb.create_sheet(title="B2B_Table_4A")
    ws_b2b.views.sheetView[0].showGridLines = True

    b2b_headers = [
        "Source Invoice ID", "Recipient GSTIN", "Recipient Name", "Invoice Number",
        "Invoice Date", "Invoice Value (₹)", "Place of Supply", "Reverse Charge",
        "Invoice Type", "Taxable Value (₹)", "Rate (%)", "CGST (₹)", "SGST (₹)", "IGST (₹)"
    ]

    for col_idx, h_text in enumerate(b2b_headers, start=1):
        cell = ws_b2b.cell(row=1, column=col_idx, value=h_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, rec in enumerate(draft.b2b_records, start=2):
        ws_b2b.cell(row=row_idx, column=1, value=rec.source_invoice_id).font = REGULAR_FONT
        ws_b2b.cell(row=row_idx, column=2, value=rec.recipient_gstin).font = REGULAR_FONT
        ws_b2b.cell(row=row_idx, column=3, value=rec.recipient_name or "").font = REGULAR_FONT
        ws_b2b.cell(row=row_idx, column=4, value=rec.invoice_number).font = REGULAR_FONT
        ws_b2b.cell(row=row_idx, column=5, value=rec.invoice_date).font = REGULAR_FONT
        
        c_inv = ws_b2b.cell(row=row_idx, column=6, value=float(rec.invoice_value))
        c_inv.number_format = NUMBER_FORMAT
        c_inv.font = REGULAR_FONT

        ws_b2b.cell(row=row_idx, column=7, value=rec.place_of_supply).font = REGULAR_FONT
        ws_b2b.cell(row=row_idx, column=8, value=rec.reverse_charge).font = REGULAR_FONT
        ws_b2b.cell(row=row_idx, column=9, value=rec.invoice_type).font = REGULAR_FONT

        c_taxable = ws_b2b.cell(row=row_idx, column=10, value=float(rec.taxable_value))
        c_taxable.number_format = NUMBER_FORMAT
        c_taxable.font = REGULAR_FONT

        c_rate = ws_b2b.cell(row=row_idx, column=11, value=float(rec.rate))
        c_rate.font = REGULAR_FONT

        c_cgst = ws_b2b.cell(row=row_idx, column=12, value=float(rec.cgst_amount))
        c_cgst.number_format = NUMBER_FORMAT
        c_cgst.font = REGULAR_FONT

        c_sgst = ws_b2b.cell(row=row_idx, column=13, value=float(rec.sgst_amount))
        c_sgst.number_format = NUMBER_FORMAT
        c_sgst.font = REGULAR_FONT

        c_igst = ws_b2b.cell(row=row_idx, column=14, value=float(rec.igst_amount))
        c_igst.number_format = NUMBER_FORMAT
        c_igst.font = REGULAR_FONT

    auto_fit_columns(ws_b2b)

    # -------------------------------------------------------------
    # 3. B2C_TABLE_7 SHEET
    # -------------------------------------------------------------
    ws_b2c = wb.create_sheet(title="B2C_Table_7")
    ws_b2c.views.sheetView[0].showGridLines = True

    b2c_headers = [
        "Source Invoice ID", "Place of Supply", "Invoice Number", "Invoice Date",
        "Invoice Value (₹)", "Taxable Value (₹)", "Rate (%)", "CGST (₹)", "SGST (₹)", "IGST (₹)"
    ]

    for col_idx, h_text in enumerate(b2c_headers, start=1):
        cell = ws_b2c.cell(row=1, column=col_idx, value=h_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, rec in enumerate(draft.b2c_records, start=2):
        ws_b2c.cell(row=row_idx, column=1, value=rec.source_invoice_id).font = REGULAR_FONT
        ws_b2c.cell(row=row_idx, column=2, value=rec.place_of_supply).font = REGULAR_FONT
        ws_b2c.cell(row=row_idx, column=3, value=rec.invoice_number or "").font = REGULAR_FONT
        ws_b2c.cell(row=row_idx, column=4, value=rec.invoice_date or "").font = REGULAR_FONT

        c_inv = ws_b2c.cell(row=row_idx, column=5, value=float(rec.invoice_value))
        c_inv.number_format = NUMBER_FORMAT
        c_inv.font = REGULAR_FONT

        c_taxable = ws_b2c.cell(row=row_idx, column=6, value=float(rec.taxable_value))
        c_taxable.number_format = NUMBER_FORMAT
        c_taxable.font = REGULAR_FONT

        ws_b2c.cell(row=row_idx, column=7, value=float(rec.rate)).font = REGULAR_FONT

        c_cgst = ws_b2c.cell(row=row_idx, column=8, value=float(rec.cgst_amount))
        c_cgst.number_format = NUMBER_FORMAT
        c_cgst.font = REGULAR_FONT

        c_sgst = ws_b2c.cell(row=row_idx, column=9, value=float(rec.sgst_amount))
        c_sgst.number_format = NUMBER_FORMAT
        c_sgst.font = REGULAR_FONT

        c_igst = ws_b2c.cell(row=row_idx, column=10, value=float(rec.igst_amount))
        c_igst.number_format = NUMBER_FORMAT
        c_igst.font = REGULAR_FONT

    auto_fit_columns(ws_b2c)

    # -------------------------------------------------------------
    # 4. EXCLUSIONS SHEET
    # -------------------------------------------------------------
    ws_exc = wb.create_sheet(title="Exclusions")
    ws_exc.views.sheetView[0].showGridLines = True

    exc_headers = [
        "Invoice ID", "Exclusion Reason", "Confidence Score", "Confidence Tier",
        "Recommendation", "Review Status", "Validation Status", "Errors", "Warnings", "Triggered Rules"
    ]

    for col_idx, h_text in enumerate(exc_headers, start=1):
        cell = ws_exc.cell(row=1, column=col_idx, value=h_text)
        cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, exc in enumerate(draft.exclusions, start=2):
        ws_exc.cell(row=row_idx, column=1, value=exc.invoice_id).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=2, value=exc.reason).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=3, value=float(exc.confidence_score)).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=4, value=exc.confidence_tier).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=5, value=exc.recommendation).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=6, value=exc.review_status).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=7, value=exc.validation_status).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=8, value=exc.error_count).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=9, value=exc.warning_count).font = REGULAR_FONT
        ws_exc.cell(row=row_idx, column=10, value="; ".join(exc.triggered_rules)).font = REGULAR_FONT

    auto_fit_columns(ws_exc)

    # -------------------------------------------------------------
    # 5. VALIDATION REPORT SHEET
    # -------------------------------------------------------------
    ws_val = wb.create_sheet(title="Validation_Report")
    ws_val.views.sheetView[0].showGridLines = True

    val_headers = ["Invoice ID", "Rule ID", "Severity", "Field", "Message"]
    for col_idx, h_text in enumerate(val_headers, start=1):
        cell = ws_val.cell(row=1, column=col_idx, value=h_text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if draft.validation_issues:
        for row_idx, iss in enumerate(draft.validation_issues, start=2):
            ws_val.cell(row=row_idx, column=1, value=iss.invoice_id).font = REGULAR_FONT
            ws_val.cell(row=row_idx, column=2, value=iss.rule_id).font = REGULAR_FONT
            ws_val.cell(row=row_idx, column=3, value=iss.severity).font = REGULAR_FONT
            ws_val.cell(row=row_idx, column=4, value=iss.field).font = REGULAR_FONT
            ws_val.cell(row=row_idx, column=5, value=iss.message).font = REGULAR_FONT
    else:
        ws_val.cell(row=2, column=1, value="ALL_INVOICES").font = REGULAR_FONT
        ws_val.cell(row=2, column=2, value="GSTR1-VALID").font = REGULAR_FONT
        ws_val.cell(row=2, column=3, value="INFO").font = REGULAR_FONT
        ws_val.cell(row=2, column=4, value="GSTR-1 Draft").font = REGULAR_FONT
        ws_val.cell(row=2, column=5, value="Zero draft validation errors detected across all admitted records.").font = REGULAR_FONT

    auto_fit_columns(ws_val)

    # Save Workbook
    wb.save(path)
    return path
